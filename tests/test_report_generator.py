import openpyxl
import pytest

from src.models import TestResult as TR
from src.report_generator import generate_report


def _result(api="https://api.example.com", test_id="VAL-001", status_code=200,
            test_type="Vulnerability", response="hello") -> TR:
    return TR(
        api=api,
        key="username",
        test_case="Some description",
        test_case_id=test_id,
        test_case_type=test_type,
        status="Accepted (200)",
        mandatory="Yes",
        status_code=status_code,
        response=response,
    )


def test_empty_results_does_not_crash(tmp_path):
    out = tmp_path / "empty.xlsx"
    generate_report({}, out)  # must not raise
    wb = openpyxl.load_workbook(out)
    assert "No Results" in wb.sheetnames


def test_formula_injection_escaped(tmp_path):
    out = tmp_path / "inj.xlsx"
    results = {
        "https://api.example.com": [
            _result(response="=HYPERLINK(\"http://evil\",\"click\")"),
            _result(response="+1+1"),
            _result(response="-1"),
            _result(response="@cmd"),
            _result(response="plain text"),
        ]
    }
    generate_report(results, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["https___api.example.com"]
    responses = [ws.cell(row=r, column=9).value for r in range(2, 7)]
    assert responses[0].startswith("'=")
    assert responses[1].startswith("'+")
    assert responses[2].startswith("'-")
    assert responses[3].startswith("'@")
    assert responses[4] == "plain text"


def test_sheet_name_collision_deduped(tmp_path):
    out = tmp_path / "collision.xlsx"
    results = {
        "https://a.example.com/endpoint-with-a-long-path-that-overflows": [_result()],
        "https://b.example.com/endpoint-with-a-long-path-that-overflows": [_result()],
    }
    generate_report(results, out)  # must not raise duplicate-sheet ValueError
    wb = openpyxl.load_workbook(out)
    assert len(wb.sheetnames) == 2


def test_empty_api_label_gets_default_sheet_name(tmp_path):
    out = tmp_path / "empty_api.xlsx"
    generate_report({"": [_result()]}, out)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames[0] != ""


def test_5xx_not_highlighted_green(tmp_path):
    out = tmp_path / "fivexx.xlsx"
    results = {"https://api.example.com": [_result(status_code=500)]}
    generate_report(results, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["https___api.example.com"]
    cell = ws.cell(row=2, column=1)
    # 5xx must not be painted green (it's a server error, not a security pass)
    assert cell.fill.start_color.rgb not in ("00C6EFCE", "FFC6EFCE")


def test_4xx_highlighted_green(tmp_path):
    out = tmp_path / "fourxx.xlsx"
    results = {"https://api.example.com": [_result(status_code=400)]}
    generate_report(results, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["https___api.example.com"]
    cell = ws.cell(row=2, column=1)
    assert cell.fill.start_color.rgb in ("00C6EFCE", "FFC6EFCE")


def test_2xx_vuln_highlighted_red(tmp_path):
    out = tmp_path / "accepted.xlsx"
    results = {"https://api.example.com": [_result(status_code=200)]}
    generate_report(results, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["https___api.example.com"]
    cell = ws.cell(row=2, column=1)
    assert cell.fill.start_color.rgb in ("00FFC7CE", "FFFFC7CE")


def test_report_has_test_case_id_column(tmp_path):
    out = tmp_path / "ids.xlsx"
    results = {"https://api.example.com": [_result(test_id="HDR-001")]}
    generate_report(results, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["https___api.example.com"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, 13)]
    assert "TestCase ID" in headers
    assert ws.cell(row=2, column=2).value == "HDR-001"