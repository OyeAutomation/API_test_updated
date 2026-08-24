#!/usr/bin/env python3
"""Quick integration test - starts server in a thread, runs tests locally."""
import sys
import os
import threading
import time

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Start Flask in a daemon thread
from mock_server.app import app

def run_server():
    app.run(port=5000, debug=False, use_reloader=False)

server_thread = threading.Thread(target=run_server, daemon=True)
server_thread.start()
time.sleep(2)

import requests

# Health check
r = requests.get("http://localhost:5000/api/protected", timeout=5)
print(f"Server ready: {r.status_code}")

# ----- Test 1: register endpoint -----
print("\n=== Test: /api/register ===")
from src.test_engine import run_suite
from src.models import Settings

curl = (
    'curl -X POST http://localhost:5000/api/register '
    '-H "Content-Type: application/json" '
    '-H "Authorization: Bearer valid-token-12345" '
    '-d \'{"username":"john_doe","email":"john@test.com","password":"secret123"}\''
)

settings = Settings(request_delay_ms=5, timeout_seconds=10, max_retries=1)

from src.api_profile import load_profiles, find_profile
from src.curl_parser import parse_curl

PROFILE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "config", "api_profiles")
PROFILES = load_profiles(PROFILE_DIR)


def _profile_for(curl: str):
    return find_profile(PROFILES, parse_curl(curl))


def _vuln_accepted(results) -> int:
    """Vulnerability rows the engine itself flagged as accepted (VULNERABLE)."""
    return sum(
        1 for r in results
        if r.test_case_type == "Vulnerability" and "VULNERABLE" in r.status
    )


results = run_suite(curl, crypto_profile_config=None, test_case_file=None,
                    settings=settings, profile=_profile_for(curl))

total = len(results)
vuln_accepted = _vuln_accepted(results)
errors = sum(1 for r in results if r.test_case_type == "Error")
print(f"Results: {total} total, {vuln_accepted} vuln accepted, {errors} errors")

# Show header results
for r in results:
    if r.test_case_id in ("HDR-001", "HDR-002", "HDR-003", "HDR-004", "HDR-005", "GLB-001", "GLB-002"):
        print(f"  {r.test_case_id} ({r.test_case_type}): {r.status} [{r.status_code}]")

# Show some payload results
for r in results:
    if r.test_case_id in ("VAL-001", "VAL-007", "VAL-009", "VAL-012"):
        print(f"  {r.test_case_id} key={r.key}: {r.status} [{r.status_code}] mandatory={r.mandatory}")

# ----- Test 2: protected endpoint -----
print("\n=== Test: /api/protected ===")
curl2 = (
    'curl -X GET http://localhost:5000/api/protected '
    '-H "Authorization: Bearer valid-token-12345"'
)
results2 = run_suite(curl2, crypto_profile_config=None, test_case_file=None,
                     settings=settings, profile=_profile_for(curl2))
print(f"Results: {len(results2)} total")

# ----- Test 3: nested endpoint -----
print("\n=== Test: /api/register/nested ===")
curl3 = (
    'curl -X POST http://localhost:5000/api/register/nested '
    '-H "Content-Type: application/json" '
    '-H "Authorization: Bearer valid-token-12345" '
    '-d \'{"user":{"name":"john","email":"john@test.com","address":{"zip":"12345"}}}\''
)
results3 = run_suite(curl3, crypto_profile_config=None, test_case_file=None,
                     settings=settings, profile=_profile_for(curl3))
print(f"Results: {len(results3)} total")

# ----- Report -----
print("\n=== Generating Report ===")
from src.report_generator import generate_report

grouped = {}
for r in results + results2 + results3:
    grouped.setdefault(r.api, []).append(r)
report_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "reports", "integration_report.xlsx")
generate_report(grouped, report_path)
print(f"Report: {report_path}")

# Verify
import openpyxl
wb = openpyxl.load_workbook(report_path)
ws = wb.active
headers = [cell.value for cell in ws[1]]
expected = ["API", "TestCase ID", "Key", "TestCase", "TestCaseType", "Status", "Mandatory", "StatusCode",
            "Response", "Decrypted Response", "Decrypted Payload Sent", "Encrypted Payload"]
assert headers == expected, f"Header mismatch: {headers}"
print(f"Verified: {ws.max_row - 1} rows, correct headers ✓")

all_results = results + results2 + results3
total_all = len(all_results)
vuln_all = _vuln_accepted(all_results)
err_all = sum(1 for r in all_results if r.test_case_type == "Error")

# Check vulnerability highlighting is meaningful: find rows where Vuln accepted
print(f"\n=== SUMMARY ===")
print(f"Total: {total_all} | Vulnerabilities accepted: {vuln_all} | Errors: {err_all}")

assert err_all == 0, f"{err_all} errors in results - integration test FAILED"
assert vuln_all == 0, f"{vuln_all} vulnerabilities accepted against the mock server - integration test FAILED"
print(f"PASSED ✓")
