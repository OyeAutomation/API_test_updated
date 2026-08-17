#!/usr/bin/env python3
"""Integration test that starts the mock server and runs the full test suite."""
import sys
import os
import json
import subprocess
import time
import threading

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import requests
from src.cli import _load_crypto_profile
from src.test_engine import run_suite
from src.report_generator import generate_report
from src.models import TestResult, Settings


def test_register_endpoint():
    print("=" * 60)
    print("Integration Test: /api/register")
    print("=" * 60)

    curl = (
        'curl -X POST http://localhost:5000/api/register '
        '-H "Content-Type: application/json" '
        '-H "Authorization: Bearer valid-token-12345" '
        '-d \'{"username":"john_doe","email":"john@test.com","password":"secret123"}\''
    )

    settings = Settings(request_delay_ms=10, timeout_seconds=10, max_retries=1)
    results = run_suite(curl, crypto_profile_config=None, test_case_file=None, settings=settings)

    print(f"\nTotal results: {len(results)}")
    vuln_accepted = sum(
        1 for r in results
        if r.test_case_type == "Vulnerability" and 0 < r.status_code < 400
    )
    errors = sum(1 for r in results if r.test_case_type == "Error")
    print(f"Vulnerabilities accepted: {vuln_accepted}")
    print(f"Errors: {errors}")

    # Generate report
    grouped = {}
    for r in results:
        grouped.setdefault(r.api, []).append(r)
    report_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "reports", "integration_report.xlsx")
    generate_report(grouped, report_path)
    print(f"Report: {report_path}")

    return results


def test_protected_endpoint():
    print("\n" + "=" * 60)
    print("Integration Test: /api/protected")
    print("=" * 60)

    curl = (
        'curl -X GET http://localhost:5000/api/protected '
        '-H "Authorization: Bearer valid-token-12345"'
    )

    settings = Settings(request_delay_ms=10, timeout_seconds=10, max_retries=1)
    results = run_suite(curl, crypto_profile_config=None, test_case_file=None, settings=settings)

    print(f"\nTotal results: {len(results)}")
    return results


def test_nested_endpoint():
    print("\n" + "=" * 60)
    print("Integration Test: /api/register/nested")
    print("=" * 60)

    curl = (
        'curl -X POST http://localhost:5000/api/register/nested '
        '-H "Content-Type: application/json" '
        '-H "Authorization: Bearer valid-token-12345" '
        '-d \'{"user":{"name":"john","email":"john@test.com","address":{"zip":"12345"}}}\''
    )

    settings = Settings(request_delay_ms=10, timeout_seconds=10, max_retries=1)
    results = run_suite(curl, crypto_profile_config=None, test_case_file=None, settings=settings)

    print(f"\nTotal results: {len(results)}")
    return results


def verify_report(report_path):
    import openpyxl
    wb = openpyxl.load_workbook(report_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    expected = ["API", "Key", "TestCase", "TestCaseType", "Status", "Mandatory", "StatusCode",
                "Response", "Decrypted Response", "Decrypted Payload Sent", "Encrypted Payload"]
    assert headers == expected, f"Header mismatch: {headers}"
    row_count = ws.max_row - 1
    print(f"\nReport verification: {row_count} data rows, correct headers ✓")
    return True


if __name__ == "__main__":
    server_proc = None
    try:
        print("Starting mock server...")
        server_proc = subprocess.Popen(
            [sys.executable, "-u", "mock_server/app.py"],
            cwd=os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )
        time.sleep(3)

        # Health check
        for attempt in range(10):
            try:
                r = requests.get("http://localhost:5000/api/protected", timeout=2)
                print("Mock server is ready!")
                break
            except requests.ConnectionError:
                time.sleep(1)
        else:
            print("Failed to connect to mock server")
            sys.exit(1)

        # Run tests
        all_results = []
        r1 = test_register_endpoint()
        all_results.extend(r1)
        r2 = test_protected_endpoint()
        all_results.extend(r2)
        r3 = test_nested_endpoint()
        all_results.extend(r3)

        # Generate combined report
        grouped = {}
        for r in all_results:
            grouped.setdefault(r.api, []).append(r)
        report_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "reports", "integration_report.xlsx"
        )
        generate_report(grouped, report_path)
        print(f"\n{'=' * 60}")
        print(f"Combined report: {report_path}")

        # Verify
        verify_report(report_path)

        # Summary
        total = len(all_results)
        vuln_accepted = sum(
            1 for r in all_results
            if r.test_case_type == "Vulnerability" and 0 < r.status_code < 400
        )
        errors = sum(1 for r in all_results if r.test_case_type == "Error")
        print(f"\nSummary: {total} results, {vuln_accepted} accepted, {errors} errors")
        print("Integration test PASSED ✓")

    finally:
        if server_proc:
            server_proc.terminate()
            server_proc.wait(timeout=5)
